from pathlib import Path

import pdfplumber
import pytesseract
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".webp"}
TEXT_EXTENSIONS = {
    ".txt",
    ".csv",
    ".tsv",
    ".md",
    ".log",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".yml",
    ".yaml",
    ".ini",
    ".cfg",
    ".sql",
    ".py",
    ".js",
    ".ts",
    ".css",
    ".sh",
    ".bat",
    ".ps1",
    ".rtf",
}


def extract_text_from_file(path: str) -> str:
    ext = Path(path).suffix.lower()

    if ext == ".pdf":
        return extract_from_pdf(path)
    if ext in IMAGE_EXTENSIONS:
        return extract_from_image(path)
    if ext == ".docx":
        return extract_from_docx(path)
    if ext in {".xlsx", ".xlsm"}:
        return extract_from_xlsx(path)
    if ext == ".xls":
        return extract_from_xls(path)
    if ext in {".pptx", ".pptm"}:
        return extract_from_pptx(path)
    if ext in TEXT_EXTENSIONS:
        return extract_from_text_like(path)

    return ""


def extract_from_pdf(path: str) -> str:
    text = ""

    try:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception:
        return ""

    if text.strip():
        return text

    try:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                img = page.to_image(resolution=300).original
                text += pytesseract.image_to_string(img, lang="rus+kaz")
    except Exception:
        return text

    return text


def extract_from_image(path: str) -> str:
    img = Image.open(path)
    try:
        return pytesseract.image_to_string(img, lang="rus+kaz")
    except Exception:
        return ""


def extract_from_docx(path: str) -> str:
    try:
        from docx import Document
    except Exception:
        return ""

    doc = Document(path)
    parts = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    for table in doc.tables:
        for row in table.rows:
            row_values = [
                cell.text.strip()
                for cell in row.cells
                if cell.text and cell.text.strip()
            ]
            if row_values:
                parts.append(" | ".join(row_values))

    return "\n".join(parts)


def extract_from_xlsx(path: str) -> str:
    wb = load_workbook(path, data_only=True)
    text_parts = []

    for sheet in wb.worksheets:
        text_parts.append(f"[{sheet.title}]")
        for row in sheet.iter_rows():
            values = [str(cell.value).strip() for cell in row if cell.value not in (None, "")]
            if values:
                text_parts.append(" | ".join(values))

    return "\n".join(text_parts)


def extract_from_xls(path: str) -> str:
    try:
        import xlrd
    except Exception:
        return ""

    book = xlrd.open_workbook(path)
    text_parts = []

    for sheet in book.sheets():
        text_parts.append(f"[{sheet.name}]")
        for row_idx in range(sheet.nrows):
            values = []
            for col_idx in range(sheet.ncols):
                value = sheet.cell_value(row_idx, col_idx)
                if value not in ("", None):
                    values.append(str(value).strip())
            if values:
                text_parts.append(" | ".join(values))

    return "\n".join(text_parts)


def extract_from_pptx(path: str) -> str:
    prs = Presentation(path)
    parts = []

    for slide in prs.slides:
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and text.strip():
                parts.append(text.strip())

    return "\n".join(parts)


def extract_from_text_like(path: str) -> str:
    raw = Path(path).read_bytes()

    for encoding in ("utf-8", "utf-8-sig", "cp1251", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue

    return ""
